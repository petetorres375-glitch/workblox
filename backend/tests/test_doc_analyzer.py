import io
from unittest.mock import patch

import pytest

from app.services.claude_client import MAX_IMAGES

MOCK_RESPONSE = {
    "summary": "A test document.",
    "key_data_points": ["Item 1"],
    "action_items": ["None identified."],
    "red_flags": ["None identified."],
}


@pytest.fixture(autouse=True)
def _bypass_tool_entitlement():
    # TESTING mode skips auth entirely, so g.user is never set — require_tool()
    # would blow up reading it. These tests are about doc-analyzer logic, not
    # access control, so bypass require_tool() and require_personal().
    with patch("app.routes.doc_analyzer.require_tool", return_value=None), \
         patch("app.routes.doc_analyzer.require_personal", return_value=None):
        yield


def test_doc_analyzer_txt(client):
    data = {"file": (io.BytesIO(b"This is a test document."), "test.txt")}
    with patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE):
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200
    result = rv.get_json()
    assert "summary" in result
    assert "key_data_points" in result
    assert "action_items" in result
    assert "red_flags" in result


def test_doc_analyzer_md(client):
    content = b"# Heading\n\nSome markdown content."
    data = {"file": (io.BytesIO(content), "notes.md")}
    with patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE):
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200


def test_doc_analyzer_docx(client):
    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Quarterly vendor agreement.")
    doc.save(buf)
    data = {"file": (io.BytesIO(buf.getvalue()), "agreement.docx")}
    with patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE) as mock_call:
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200
    assert "Quarterly vendor agreement." in mock_call.call_args.kwargs["user_message"]


def test_doc_analyzer_no_file(client):
    rv = client.post("/api/doc", data={}, content_type="multipart/form-data")
    assert rv.status_code == 400


def test_doc_analyzer_unsupported_type(client):
    data = {"file": (io.BytesIO(b"data"), "image.jpg")}
    rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 415


def test_doc_analyzer_empty_file(client):
    data = {"file": (io.BytesIO(b"   "), "blank.txt")}
    rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 422


def test_doc_analyzer_photo_single(client):
    data = {"images": (io.BytesIO(b"fake-jpeg-bytes"), "photo-1.jpg")}
    with patch("app.routes.doc_analyzer.prepare_image", return_value=b"normalized-jpeg-bytes"), \
         patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE) as mock_call:
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200
    assert rv.get_json()["summary"] == MOCK_RESPONSE["summary"]
    _, kwargs = mock_call.call_args
    assert len(kwargs["images"]) == 1
    assert kwargs["images"][0]["media_type"] == "image/jpeg"


def test_doc_analyzer_photo_multiple_pages(client):
    data = {"images": [
        (io.BytesIO(b"page-1"), "photo-1.jpg"),
        (io.BytesIO(b"page-2"), "photo-2.jpg"),
        (io.BytesIO(b"page-3"), "photo-3.jpg"),
    ]}
    with patch("app.routes.doc_analyzer.prepare_image", return_value=b"normalized-jpeg-bytes"), \
         patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE) as mock_call:
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200
    _, kwargs = mock_call.call_args
    assert len(kwargs["images"]) == 3


def test_doc_analyzer_photo_and_file_rejected(client):
    data = {
        "file": (io.BytesIO(b"This is a test document."), "test.txt"),
        "images": (io.BytesIO(b"fake-jpeg-bytes"), "photo-1.jpg"),
    }
    rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 400


def test_doc_analyzer_photo_too_many(client):
    data = {"images": [
        (io.BytesIO(f"page-{i}".encode()), f"photo-{i}.jpg") for i in range(MAX_IMAGES + 1)
    ]}
    rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 400


def test_doc_analyzer_photo_unreadable(client):
    data = {"images": (io.BytesIO(b"not-an-image"), "photo-1.jpg")}
    with patch("app.routes.doc_analyzer.prepare_image",
               side_effect=ValueError("Could not read 'photo-1.jpg' as an image.")):
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 415


def test_doc_email_no_key(client):
    # The route reads SENDGRID_API_KEY straight from the environment, and
    # Config.load_dotenv() has already pulled the real key out of backend/.env
    # by the time tests run -- so without forcing it empty here, this test
    # can't reach the branch it exists to cover, falls through to a live
    # SendGrid request, and fails with a 500 instead. patch.dict restores the
    # original value afterwards, matching test_doc_email_sends below.
    import os
    payload = {"email": "test@example.com", "filename": "doc.txt", "result": MOCK_RESPONSE}
    with patch.dict(os.environ, {"SENDGRID_API_KEY": ""}):
        rv = client.post("/api/doc/email", json=payload)
    assert rv.status_code == 503
    assert "not configured" in rv.get_json()["error"]


def test_doc_email_invalid_email(client):
    payload = {"email": "notanemail", "filename": "doc.txt", "result": MOCK_RESPONSE}
    rv = client.post("/api/doc/email", json=payload)
    assert rv.status_code == 400


def test_doc_email_missing_result(client):
    payload = {"email": "test@example.com", "filename": "doc.txt"}
    rv = client.post("/api/doc/email", json=payload)
    assert rv.status_code == 400


def test_doc_email_sends(client):
    import os
    from unittest.mock import MagicMock, patch
    payload = {"email": "pete@example.com", "filename": "report.txt", "result": MOCK_RESPONSE}
    mock_resp = MagicMock()
    mock_resp.__enter__ = lambda s: s
    mock_resp.__exit__ = MagicMock(return_value=False)
    mock_resp.read.return_value = b""
    with patch.dict(os.environ, {"SENDGRID_API_KEY": "SG.test"}), \
         patch("urllib.request.urlopen", return_value=mock_resp):
        rv = client.post("/api/doc/email", json=payload)
    assert rv.status_code == 200
    assert rv.get_json()["success"] is True


# ── Scanned-page limit ───────────────────────────────────────────────────────

def _scanned_pdf(pages):
    """Image-only pages (no text layer), like a phone scan."""
    import fitz
    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page()
        pix = fitz.Pixmap(fitz.csGRAY, fitz.IRect(0, 0, 60, 80), False)
        pix.clear_with(200)
        page.insert_image(page.rect, pixmap=pix)
    return doc.tobytes()


def _typed_pdf(pages):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_font("Helvetica", size=10)
    for p in range(pages):
        pdf.add_page()
        pdf.cell(0, 5, f"Typed page {p + 1}")
    return bytes(pdf.output())


def test_doc_analyzer_rejects_a_long_scan_before_any_ocr(client):
    from app.routes.doc_analyzer import MAX_SCANNED_PAGES
    data = {"file": (io.BytesIO(_scanned_pdf(MAX_SCANNED_PAGES + 1)), "scan.pdf")}
    with patch("app.services.file_handler.ocr_png") as ocr, \
         patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE) as ai:
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 413
    assert rv.get_json()["code"] == "too_many_scanned_pages"
    assert f"{MAX_SCANNED_PAGES + 1} scanned pages" in rv.get_json()["error"]
    ocr.assert_not_called()
    ai.assert_not_called()


def test_doc_analyzer_accepts_a_scan_at_the_limit(client):
    from app.routes.doc_analyzer import MAX_SCANNED_PAGES
    data = {"file": (io.BytesIO(_scanned_pdf(MAX_SCANNED_PAGES)), "scan.pdf")}
    with patch("app.services.file_handler.ocr_png", return_value="scanned words"), \
         patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE):
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200


def test_doc_analyzer_has_no_page_limit_for_typed_pdfs(client):
    data = {"file": (io.BytesIO(_typed_pdf(40)), "long.pdf")}
    with patch("app.routes.doc_analyzer.claude_client.call", return_value=MOCK_RESPONSE):
        rv = client.post("/api/doc", data=data, content_type="multipart/form-data")
    assert rv.status_code == 200
